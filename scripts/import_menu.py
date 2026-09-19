#!/usr/bin/env python3
"""Import the dining hall Excel sheets into the static menu catalogue."""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from tagging import apply_tags


MERCHANT_NAMES = {
    "中式点心": "中式点心",
    "民族风味": "民族风味餐厅",
    "淮扬快餐": "淮扬餐厅",
    "自选教工": "教工自选",
    "超霖美食": "超霖美食",
}


def clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_workbooks(source_dir: Path) -> dict[str, Any]:
    items: list[dict[str, Any]] = []
    for workbook_path in sorted(source_dir.rglob("*.xlsx")):
        merchant = next(
            (label for part, label in MERCHANT_NAMES.items() if part in workbook_path.name),
            workbook_path.stem,
        )
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        for sheet in workbook.worksheets:
            location: str | None = None
            stall: str | None = None
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=4, max_col=6, values_only=True), start=4
            ):
                serial, row_location, row_stall, dish_name, raw_price, raw_note = row
                if clean(row_location):
                    location = clean(row_location)
                if clean(row_stall):
                    stall = clean(row_stall)
                name = clean(dish_name)
                if not name or not isinstance(serial, (int, float)):
                    continue

                record: dict[str, Any] = {
                    "id": f"{workbook_path.stem}:{sheet.title}:{row_number}",
                    "location": location or "第一餐饮大楼",
                    "merchant": merchant,
                    "stall": stall or merchant,
                    "category": clean(sheet.title) or "其他",
                    "name": name,
                    "price": None,
                    "priceLabel": None,
                    "note": clean(raw_note),
                    "source": workbook_path.name,
                    "sourceSheet": sheet.title,
                }
                if isinstance(raw_price, (int, float)) and not isinstance(raw_price, bool):
                    record["price"] = raw_price
                    if raw_price == 0:
                        record["priceNote"] = "原表价格为 0，含义待核实"
                elif clean(raw_price):
                    record["priceLabel"] = clean(raw_price)
                items.append(record)
        workbook.close()

    overrides_path = Path(__file__).with_name("manual_tag_overrides.json")
    overrides = json.loads(overrides_path.read_text(encoding="utf-8")) if overrides_path.exists() else {}
    for item in items:
        apply_tags(item, overrides)

    numeric_prices = [item for item in items if item["price"] is not None]
    summary = {
        "itemCount": len(items),
        "numericPriceCount": len(numeric_prices),
        "missingPriceCount": sum(
            1 for item in items if item["price"] is None and item["priceLabel"] is None
        ),
        "priceLabelCount": sum(1 for item in items if item["priceLabel"] is not None),
        "zeroPriceCount": sum(1 for item in numeric_prices if item["price"] == 0),
        "noteCount": sum(1 for item in items if item["note"]),
        "merchantCount": len({item["merchant"] for item in items}),
        "categoryCount": len({item["category"] for item in items}),
        "tasteLabeledCount": sum(1 for item in items if item["tasteTags"]),
        "tastePendingCount": sum(1 for item in items if not item["tasteTags"]),
        "formTaggedCount": sum(1 for item in items if item["formTags"] != ["其他"]),
        "mealTaggedCount": sum(1 for item in items if item["mealTags"]),
        "mealPendingCount": sum(1 for item in items if not item["mealTags"]),
    }
    return {
        "venue": "第一餐饮大楼",
        "importedAt": datetime.now().astimezone().isoformat(timespec="minutes"),
        "sourceFiles": len(list(source_dir.rglob("*.xlsx"))),
        "tagVersion": "keyword-rules-v2-meals",
        "summary": summary,
        "items": items,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Import a folder of dining menu Excel files")
    parser.add_argument("source_dir", type=Path, help="Folder containing the .xlsx files")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("dist/data/menu.json"),
        help="Output JSON path (default: dist/data/menu.json)",
    )
    args = parser.parse_args()
    if not args.source_dir.is_dir():
        parser.error(f"Source directory does not exist: {args.source_dir}")
    data = import_workbooks(args.source_dir)
    if not data["items"]:
        parser.error("No menu rows found. Check the workbook names and columns A–F.")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(data["summary"], ensure_ascii=False))
    print(f"Wrote {len(data['items'])} menu rows to {args.output}")


if __name__ == "__main__":
    main()
